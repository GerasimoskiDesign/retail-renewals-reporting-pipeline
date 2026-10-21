import datetime as dt
import json

import pandas as pd
from openpyxl import load_workbook

from renewals.pipeline import compute_month, processed_dir, reports_dir, run_month
from renewals.transform import metrics

MONTH = dt.date(2026, 9, 1)


def _snapshot(output_dir):
    proc = processed_dir(output_dir, MONTH)
    frames = {p.name: pd.read_parquet(p) for p in sorted(proc.glob("*.parquet"))}
    raw = {p.name: p.read_bytes() for p in sorted(proc.glob("*.parquet"))}
    audit = json.loads((proc / "audit.json").read_text())
    cells = {}
    for path in sorted(reports_dir(output_dir, MONTH).glob("*.xlsx")):
        wb = load_workbook(path)
        cells[path.name] = {
            ws.title: [[c.value for c in row] for row in ws.iter_rows()] for ws in wb.worksheets
        }
    return frames, raw, audit, cells


def test_recompute_twice_gives_identical_output(month_downloads, world, tmp_path):
    as_of = dt.date(2026, 9, 12)
    run_month(month_downloads, tmp_path, world.shops, MONTH, as_of)
    first = _snapshot(tmp_path)
    run_month(month_downloads, tmp_path, world.shops, MONTH, as_of)
    second = _snapshot(tmp_path)
    for name in first[0]:
        pd.testing.assert_frame_equal(first[0][name], second[0][name])
    assert first[1] == second[1]  # Parquet files byte-identical
    assert first[2] == second[2]
    assert first[3] == second[3]  # every Excel cell identical
    assert len(first[3]) == 12  # one report per day, 1..12 September


def test_late_completions_change_earlier_days(month_downloads):
    """Why the month is recomputed from day 1: day 8 grows as later completions arrive."""
    _, early = compute_month(month_downloads, MONTH, dt.date(2026, 9, 8))
    _, later = compute_month(month_downloads, MONTH, dt.date(2026, 9, 14))
    day8 = pd.Timestamp(2026, 9, 8)
    n_early = int((early.facts["day"] == day8).sum())
    n_later = int((later.facts["day"] == day8).sum())
    assert n_later > n_early


def test_daily_counts_never_exceed_eventual_counts_except_by_cancellation(month_downloads):
    _, early = compute_month(month_downloads, MONTH, dt.date(2026, 9, 10))
    _, final = compute_month(month_downloads, MONTH, dt.date(2026, 9, 30))
    lost = set(early.facts["process_id"]) - set(final.facts["process_id"])
    reasons = final.exclusions.set_index("process_id").loc[sorted(lost), "reason"]
    assert set(reasons) <= {"cancelled", "superseded_by_restart"}


def test_excel_report_is_formatted(month_downloads, world, tmp_path):
    as_of = dt.date(2026, 9, 3)
    outcome = run_month(month_downloads, tmp_path, world.shops, MONTH, as_of)
    path = reports_dir(tmp_path, MONTH) / "Renewals Report 03.09.2026.xlsx"
    assert path in outcome.published
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Shops", "Agents MTD"]
    shops = wb["Shops"]
    assert shops.freeze_panes == "A2"
    assert shops.auto_filter.ref.startswith("A1:")
    assert shops["A1"].font.bold and shops["A1"].fill.fgColor.rgb.endswith("1F3A5F")
    header = [c.value for c in shops[1]]
    share_col = header.index("Digital share (MTD)") + 1
    combined_col = header.index("Combined (MTD)") + 1
    assert shops.cell(2, share_col).number_format == "0.0%"
    assert shops.cell(2, combined_col).number_format == "#,##0"
    assert shops.max_row == len(world.shops) + 1

    mtd = outcome.result.facts
    mtd = mtd[mtd["day"] <= pd.Timestamp(as_of)]
    total = sum(shops.cell(r, combined_col).value for r in range(2, shops.max_row + 1))
    assert total == len(mtd)
    summary = {row[0].value: row[2].value for row in wb["Summary"].iter_rows(min_row=2, max_row=6)}
    assert summary["Combined"] == len(mtd)


def test_parquet_outputs_are_consistent(month_downloads, world, tmp_path):
    run_month(month_downloads, tmp_path, world.shops, MONTH, dt.date(2026, 9, 7))
    proc = processed_dir(tmp_path, MONTH)
    facts = pd.read_parquet(proc / "facts.parquet")
    sd = pd.read_parquet(proc / "shop_day.parquet")
    ad = pd.read_parquet(proc / "agent_day.parquet")
    assert facts["process_id"].is_unique
    assert sd["combined"].sum() == ad["combined"].sum() == len(facts)
    pd.testing.assert_frame_equal(sd, metrics.shop_day(facts), check_dtype=False)
