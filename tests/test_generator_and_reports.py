import datetime as dt

import pandas as pd
from openpyxl import load_workbook

from renewals.portal.reports import (
    CONTRACT_RENEWALS,
    CRM_EVENTS,
    DIGITAL_ORDERS,
    STORE_COMPLETIONS,
    TARIFF_MIGRATIONS,
    WEB_ORDERS,
    build_report,
    report_to_xlsx,
)
from renewals.transform.sources import read_report
from renewals.world.generator import generate_world
from renewals.world.model import CLICK_COLLECT, IN_STORE, WorldConfig

D = dt.date


def test_same_seed_same_world_different_seed_different_world(world):
    again = generate_world(WorldConfig(seed=42))
    pd.testing.assert_frame_equal(world.processes, again.processes)
    pd.testing.assert_frame_equal(world.events, again.events)
    other = generate_world(WorldConfig(seed=7))
    assert not world.processes["contract_id"].equals(other.processes["contract_id"])


def test_world_size_matches_spec(world):
    assert world.shops["region_id"].nunique() == 4
    assert 25 <= len(world.shops) <= 35
    assert 150 <= len(world.agents) <= 250


def test_phone_numbers_use_fake_prefix(world):
    assert world.processes["msisdn"].str.startswith("+999").all()


def test_completion_happens_zero_to_six_days_after_start(world):
    p = world.processes.dropna(subset=["completed_at"])
    lag = (p["completed_at"].dt.normalize() - p["started_at"].dt.normalize()).dt.days
    assert lag.between(0, 6).all()


def test_each_contract_is_one_sale_apart_from_injected_traps(world):
    p = world.processes
    plain = p[~p["is_companion"] & p["restart_of"].isna()]
    assert plain["contract_id"].is_unique


def test_all_four_traps_are_present(world):
    p = world.processes
    assert p["is_companion"].sum() > 100  # trap 1
    assert (p["channel"] == CLICK_COLLECT).sum() > 1000  # trap 2
    prior = (p["started_at"].dt.month == 8) & (p["completed_at"].dt.month == 9)
    assert prior.sum() > 50  # trap 3
    assert p["restart_of"].notna().sum() > 100  # trap 4
    restarted = p.loc[p["restart_of"].notna(), "restart_of"]
    unlogged = p[p["process_id"].isin(restarted) & ~p["cancel_logged"]]
    assert len(unlogged) > 0  # some restarts lack a logged cancellation


def test_campaign_week_has_about_three_times_normal_volume(world):
    p = world.processes
    days = p["started_at"].dt.normalize()
    weekday = p["started_at"].dt.weekday < 5
    campaign = days.between(
        pd.Timestamp(world.cfg.campaign_start), pd.Timestamp(world.cfg.campaign_end)
    )
    normal = days.between(pd.Timestamp(2026, 9, 1), pd.Timestamp(2026, 9, 11))
    per_day_campaign = (campaign & weekday).sum() / 5
    per_day_normal = (normal & weekday).sum() / 9
    assert 2.5 <= per_day_campaign / per_day_normal <= 3.5


def test_shops_are_closed_on_sunday(world):
    p = world.processes
    sunday_store = p[(p["channel"] == IN_STORE) & (p["started_at"].dt.weekday == 6)]
    sunday_pickup = p[
        (p["channel"] == CLICK_COLLECT)
        & (p["completed_at"].dt.weekday == 6)
        & p["restart_of"].isna()
    ]
    assert sunday_store.empty and sunday_pickup.empty


def test_digital_orders_overlap_with_completion_reports(world):
    """Trap 1/2 at report level: every digital order is also in a completion report."""
    day = D(2026, 9, 16)
    digital = build_report(world, DIGITAL_ORDERS, day)
    comp = pd.concat(
        [build_report(world, CONTRACT_RENEWALS, day), build_report(world, TARIFF_MIGRATIONS, day)]
    )
    assert len(digital) > 0
    assert set(digital["Process Ref"]) <= set(comp["Process ID"])


def test_click_and_collect_is_in_digital_orders_and_store_pickups(world):
    day = D(2026, 9, 16)
    digital = build_report(world, DIGITAL_ORDERS, day)
    store = build_report(world, STORE_COMPLETIONS, day)
    cc = set(digital.loc[digital["Digital Channel"] == "Click & Collect", "Process Ref"])
    pickups = set(store.loc[store["Completion Type"] == "Pickup", "Process Ref"])
    assert cc and cc == pickups


def test_web_orders_day_to_next_day_files_overlap(world):
    a = build_report(world, WEB_ORDERS, D(2026, 9, 9))
    b = build_report(world, WEB_ORDERS, D(2026, 9, 10))
    shared = set(a["Order ID"]) & set(b["Order ID"])
    placed_10th = set(b.loc[b["Order Placed"].dt.date == D(2026, 9, 10), "Order ID"])
    assert shared == placed_10th and shared


def test_bi_export_has_title_rows_and_crm_export_does_not(world, tmp_path):
    day = D(2026, 9, 16)
    for spec, header_row in ((CONTRACT_RENEWALS, 4), (CRM_EVENTS, 1)):
        df = build_report(world, spec, day)
        path = tmp_path / f"{spec.key}.xlsx"
        path.write_bytes(report_to_xlsx(spec, day, df))
        ws = load_workbook(path).active
        assert ws.cell(header_row, 1).value == df.columns[0]
        back = read_report(path, spec)
        assert list(back.columns) == list(df.columns)
        assert len(back) == len(df)
