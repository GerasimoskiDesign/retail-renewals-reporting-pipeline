import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from renewals.portal.app import create_app
from renewals.settings import PortalCredentials

BI = PortalCredentials("http://testserver/bi", "bi.user", "bi-pass")
CRM = PortalCredentials("http://testserver/crm", "crm.user", "crm-pass")


@pytest.fixture()
def client(world):
    return TestClient(create_app(lambda: world, BI, CRM))


def login(client, portal, user, pw):
    return client.post(
        f"/{portal}/login", data={"username": user, "password": pw}, follow_redirects=False
    )


def test_download_requires_login(client):
    r = client.get(
        "/bi/download",
        params={"report": "digital_orders", "date_from": "2026-09-14", "date_to": ""},
    )
    assert r.status_code == 401


def test_report_page_redirects_to_login_when_signed_out(client):
    r = client.get("/bi/reports", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "login"


def test_wrong_password_is_rejected(client):
    r = login(client, "bi", "bi.user", "nope")
    assert r.status_code == 200 and "Invalid credentials" in r.text
    assert "portal_session" not in r.cookies


def test_bi_download_returns_xlsx_attachment(client):
    assert login(client, "bi", "bi.user", "bi-pass").status_code == 303
    r = client.get(
        "/bi/download",
        params={"report": "contract_renewals", "date_from": "2026-09-14", "date_to": ""},
    )
    assert r.status_code == 200
    assert "attachment" in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws["A1"].value == "Contract Renewals"
    assert ws["A2"].value == "Report date: 14.09.2026"


def test_web_orders_only_accepts_day_to_next_day_range(client):
    login(client, "bi", "bi.user", "bi-pass")
    base = {"report": "web_orders", "date_from": "2026-09-14"}
    assert client.get("/bi/download", params={**base, "date_to": ""}).status_code == 400
    assert client.get("/bi/download", params={**base, "date_to": "2026-09-16"}).status_code == 400
    ok = client.get("/bi/download", params={**base, "date_to": "2026-09-15"})
    assert ok.status_code == 200


def test_crm_names_its_own_file(client):
    login(client, "crm", "crm.user", "crm-pass")
    r = client.get(
        "/crm/download", params={"report": "crm_process_events", "date_from": "2026-09-14"}
    )
    assert r.status_code == 200
    assert 'filename="CRM_Process_Events_20260914.xlsx"' in r.headers["content-disposition"]


def test_portals_have_separate_sessions(client):
    login(client, "bi", "bi.user", "bi-pass")
    r = client.get(
        "/crm/download", params={"report": "crm_process_events", "date_from": "2026-09-14"}
    )
    assert r.status_code == 401
    # and CRM credentials do not open the BI portal
    assert "Invalid credentials" in login(client, "bi", "crm.user", "crm-pass").text


def test_bi_portal_does_not_serve_crm_report(client):
    login(client, "bi", "bi.user", "bi-pass")
    r = client.get(
        "/bi/download", params={"report": "crm_process_events", "date_from": "2026-09-14"}
    )
    assert r.status_code == 404
