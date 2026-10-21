"""Formatted daily Excel report (openpyxl).

Shop managers open this file on a phone or a small laptop, so it is optimised for that:
frozen header row, autofilter on every table, percentages as percentages and nothing
that needs a formula to recalculate.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from renewals.transform.metrics import aggregate

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
INT_FORMAT = "#,##0"
PCT_FORMAT = "0.0%"


def report_file_name(day: dt.date) -> str:
    return f"Renewals Report {day:%d.%m.%Y}.xlsx"


def _write_table(ws: Worksheet, df: pd.DataFrame, formats: dict[str, str]) -> None:
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if pd.isna(v) else v for v in row])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for idx, col in enumerate(df.columns, start=1):
        letter = get_column_letter(idx)
        fmt = formats.get(col)
        if fmt:
            for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                cell.number_format = fmt
        width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)]) + 2
        ws.column_dimensions[letter].width = min(max(width, 10), 40)
    ws.freeze_panes = "A2"
    last = get_column_letter(len(df.columns))
    ws.auto_filter.ref = f"A1:{last}{max(len(df) + 1, 1)}"


def _metric_formats(columns: list[str]) -> dict[str, str]:
    return {
        c: (PCT_FORMAT if "share" in c.lower() else INT_FORMAT)
        for c in columns
        if any(k in c.lower() for k in ("renewals", "migrations", "combined", "digital"))
    }


def _labelled(df: pd.DataFrame, shops: pd.DataFrame) -> pd.DataFrame:
    return shops[["region_name", "shop_id", "shop_name"]].merge(df, on="shop_id", how="left")


def write_daily_report(
    facts: pd.DataFrame,
    shops: pd.DataFrame,
    day: dt.date,
    as_of: dt.date,
    exclusions: pd.DataFrame,
    path: Path,
) -> Path:
    """Write the report for ``day`` (day figures + month-to-date up to ``day``)."""
    ts = pd.Timestamp(day)
    mtd = facts[facts["day"] <= ts]
    today = facts[facts["day"] == ts]

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    tot_day = aggregate(today, [])
    tot_mtd = aggregate(mtd, [])
    summary = pd.DataFrame(
        {
            "Metric": [
                "Contract renewals",
                "Tariff migrations",
                "Combined",
                "Digital",
                "Digital share",
            ],
            "Day": [
                tot_day.iloc[0][c] if len(tot_day) else 0
                for c in ("renewals", "migrations", "combined", "digital", "digital_share")
            ],
            "Month to date": [
                tot_mtd.iloc[0][c] if len(tot_mtd) else 0
                for c in ("renewals", "migrations", "combined", "digital", "digital_share")
            ],
        }
    )
    _write_table(ws, summary, {})
    for row in ws.iter_rows(min_row=2, max_row=5, min_col=2, max_col=3):
        for cell in row:
            cell.number_format = INT_FORMAT
    for cell in ws[6][1:3]:
        cell.number_format = PCT_FORMAT
    ws.append([])
    ws.append(["Report day", day.strftime("%d.%m.%Y")])
    ws.append(["Data as of", as_of.strftime("%d.%m.%Y")])
    ws.append(["Attribution", "process start day; counted once completed"])
    if not exclusions.empty:
        ws.append([])
        ws.append(["Excluded this month (all days)", None])
        for reason, n in exclusions["reason"].value_counts().sort_index().items():
            ws.append([reason, int(n)])

    shop_day = aggregate(today, ["shop_id"])
    shop_mtd = aggregate(mtd, ["shop_id"])
    names = {
        "renewals": "Renewals",
        "migrations": "Migrations",
        "combined": "Combined",
        "digital": "Digital",
        "digital_share": "Digital share",
    }
    day_part = shop_day.set_index("shop_id")[list(names)].rename(
        columns={k: f"{v} (day)" for k, v in names.items()}
    )
    mtd_part = shop_mtd.set_index("shop_id")[list(names)].rename(
        columns={k: f"{v} (MTD)" for k, v in names.items()}
    )
    shops_tbl = _labelled(day_part.join(mtd_part, how="outer").reset_index(), shops)
    shops_tbl = shops_tbl.rename(
        columns={"region_name": "Region", "shop_id": "Shop", "shop_name": "Shop name"}
    )
    for c in shops_tbl.columns[3:]:
        if "share" not in c:
            shops_tbl[c] = shops_tbl[c].fillna(0).astype(int)
    _write_table(wb.create_sheet("Shops"), shops_tbl, _metric_formats(list(shops_tbl.columns)))

    agents = _labelled(aggregate(mtd, ["shop_id", "agent_id"]), shops).dropna(subset=["agent_id"])
    agents = agents.rename(
        columns={
            "region_name": "Region",
            "shop_id": "Shop",
            "shop_name": "Shop name",
            "agent_id": "Agent",
            "renewals": "Renewals",
            "migrations": "Migrations",
            "combined": "Combined",
            "digital": "Digital",
            "digital_share": "Digital share",
        }
    )
    _write_table(wb.create_sheet("Agents MTD"), agents, _metric_formats(list(agents.columns)))

    # Fixed document timestamps keep re-runs byte-stable apart from the zip container.
    stamp = dt.datetime.combine(as_of, dt.time())
    wb.properties.created = stamp
    wb.properties.modified = stamp
    wb.properties.creator = "renewals pipeline"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
